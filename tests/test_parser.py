"""Document parser tests."""

from __future__ import annotations

import io
import unittest

from docx import Document

from parser import DocumentParseError, DocumentParser


class DocumentParserTests(unittest.TestCase):
    """Verify DOCX extraction and unsafe-type rejection."""

    def test_extracts_docx_paragraphs_and_tables(self) -> None:
        """Extract visible text from both paragraphs and tables."""
        document = Document()
        document.add_heading("Ayşe Yılmaz", level=1)
        document.add_paragraph("Python Developer olarak örnek deneyim.")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Yetenek"
        table.cell(0, 1).text = "Python"
        buffer = io.BytesIO()
        document.save(buffer)

        parsed = DocumentParser().parse_and_save("test_cv.docx", buffer.getvalue())

        self.assertIn("Ayşe Yılmaz", parsed.text)
        self.assertIn("Python", parsed.text)
        self.assertEqual(parsed.extension, "docx")

    def test_rejects_unsupported_files(self) -> None:
        """Reject extension types outside the allowlist."""
        with self.assertRaises(DocumentParseError):
            DocumentParser().parse_and_save("resume.txt", b"plain text")


if __name__ == "__main__":
    unittest.main()
