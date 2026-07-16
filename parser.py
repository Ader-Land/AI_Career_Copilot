"""Secure PDF and DOCX text extraction."""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

import fitz
import pdfplumber
from docx import Document

from config import settings
from utils import compact_text, logger, sanitize_filename, sha256_bytes


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}


class DocumentParseError(ValueError):
    """Raised when a CV cannot be safely parsed."""


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    """Extracted document and storage metadata."""

    original_filename: str
    stored_filename: str
    extension: str
    file_hash: str
    text: str
    page_count: int
    size_bytes: int


class DocumentParser:
    """Extract text from in-memory PDF and DOCX documents."""

    def parse_and_save(self, filename: str, content: bytes) -> ParsedDocument:
        """Validate, persist and extract an uploaded CV."""
        extension = Path(filename).suffix.lower()
        self._validate(filename, extension, content)
        file_hash = sha256_bytes(content)
        safe_name = sanitize_filename(filename)
        stored_filename = f"{file_hash[:16]}_{safe_name}"
        destination = settings.uploads_dir / stored_filename

        if extension == ".pdf":
            text, page_count = self._extract_pdf(content)
        else:
            text, page_count = self._extract_docx(content)

        normalized_text = compact_text(text)
        if len(normalized_text) < 20:
            raise DocumentParseError(
                "Belgeden yeterli metin çıkarılamadı. Taranmış PDF yerine metin "
                "tabanlı PDF veya DOCX yükleyin."
            )
        destination.write_bytes(content)
        logger.info(
            "document_saved hash=%s type=%s pages=%s size=%s",
            file_hash[:12],
            extension,
            page_count,
            len(content),
        )
        return ParsedDocument(
            original_filename=Path(filename).name,
            stored_filename=stored_filename,
            extension=extension.lstrip("."),
            file_hash=file_hash,
            text=normalized_text,
            page_count=page_count,
            size_bytes=len(content),
        )

    @staticmethod
    def _validate(filename: str, extension: str, content: bytes) -> None:
        """Validate extension, size and basic file signatures."""
        if extension not in SUPPORTED_EXTENSIONS:
            raise DocumentParseError("Yalnızca PDF ve DOCX dosyaları desteklenir.")
        if not content:
            raise DocumentParseError("Yüklenen dosya boş.")
        if len(content) > settings.max_upload_mb * 1024 * 1024:
            raise DocumentParseError(
                f"Dosya boyutu {settings.max_upload_mb} MB sınırını aşıyor."
            )
        if extension == ".pdf" and not content.startswith(b"%PDF"):
            raise DocumentParseError("Dosya uzantısı PDF ancak içeriği geçerli değil.")
        if extension == ".docx" and not content.startswith(b"PK"):
            raise DocumentParseError("Dosya uzantısı DOCX ancak içeriği geçerli değil.")
        if not Path(filename).name:
            raise DocumentParseError("Geçersiz dosya adı.")

    @staticmethod
    def _extract_pdf(content: bytes) -> tuple[str, int]:
        """Extract PDF text with PyMuPDF, falling back to pdfplumber."""
        try:
            with fitz.open(stream=content, filetype="pdf") as document:
                if document.needs_pass:
                    raise DocumentParseError("Parola korumalı PDF desteklenmiyor.")
                text = "\n".join(page.get_text("text") for page in document)
                page_count = document.page_count
            if text.strip():
                return text, page_count
        except DocumentParseError:
            raise
        except Exception as exc:
            logger.warning("pymupdf_extract_failed error=%s", type(exc).__name__)

        try:
            with pdfplumber.open(io.BytesIO(content)) as document:
                text = "\n".join(page.extract_text() or "" for page in document.pages)
                return text, len(document.pages)
        except Exception as exc:
            raise DocumentParseError(f"PDF okunamadı: {exc}") from exc

    @staticmethod
    def _extract_docx(content: bytes) -> tuple[str, int]:
        """Extract paragraphs and tables from a DOCX file."""
        try:
            document = Document(io.BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts), 1
        except Exception as exc:
            raise DocumentParseError(f"DOCX okunamadı: {exc}") from exc
